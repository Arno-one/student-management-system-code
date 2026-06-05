"""
成绩管理业务逻辑层 — MVC中的Model业务逻辑部分
"""
from sqlalchemy.orm import Session
from decimal import Decimal, InvalidOperation
from typing import List, Optional
from io import BytesIO
import csv
import io
import math

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from pydantic import ValidationError

from DAO import score_dao
from DAO.student_dao import get_by_student_no
from model.Score import Score
from scheme.schema_score import Addscore, Updatescore
from util.log import get_logger

# 本模块专用 logger，来源标记为 service.score_service
logger = get_logger(__name__)

# ============ 成绩 Excel/CSV 导入相关常量 ============
IMPORT_COLUMN_MAP = {
    "学号": "student_no",
    "考试次序": "exam_order",
    "成绩": "score",
}
IMPORT_HEADERS = list(IMPORT_COLUMN_MAP.keys())
IMPORT_REQUIRED_FIELDS = ["student_no", "exam_order", "score"]
FIELD_CN = {
    "student_no": "学号",
    "exam_order": "考试次序",
    "score": "成绩",
}


def add_score(data: Addscore, db: Session):
    """添加单个学生成绩（含业务校验）"""
    exist = get_by_student_no(data.student_no, db)
    if not exist:
        logger.warning("新增成绩校验失败：学生 %s 不在学生表中", data.student_no)
        raise ValueError(f"学生 {data.student_no} 的信息不在学生表中")

    is_exist = score_dao.is_exist(db, data.student_no, data.exam_order)
    if is_exist:
        logger.warning("新增成绩校验失败：学生 %s 第%s次成绩已存在",
                       data.student_no, data.exam_order)
        raise ValueError(f"学生 {data.student_no} 的第{data.exam_order}次成绩已经存在")

    new_score = Score(
        student_no=data.student_no,
        exam_order=data.exam_order,
        score=data.score
    )
    return score_dao.add_score(db, new_score)


def batch_add_scores(score_list: List[Addscore], db: Session):
    """批量添加学生成绩（含业务校验）"""
    logger.info("批量新增成绩开始：本批次 %s 条", len(score_list))

    not_exist = []
    for item in score_list:
        exist = get_by_student_no(item.student_no, db)
        if not exist:
            not_exist.append(item.student_no)
    if not_exist:
        logger.warning("批量新增成绩校验失败：以下学生不在学生表中 %s", not_exist)
        raise ValueError(f"学生 {not_exist} 的信息不在学生表中")

    batch_data = []
    for item in score_list:
        is_exist = score_dao.is_exist(db=db, student_no=item.student_no, exam_order=item.exam_order)
        if is_exist:
            logger.warning("批量新增成绩校验失败：学生 %s 第%s次成绩已存在",
                           item.student_no, item.exam_order)
            raise ValueError(f"学生 {item.student_no} 第{item.exam_order}次成绩已存在")

        batch_data.append(Score(
            student_no=item.student_no,
            exam_order=item.exam_order,
            score=item.score
        ))

    result = score_dao.batch_add_scores(db, batch_data)
    logger.info("批量新增成绩完成：成功入库 %s 条", len(result))
    return result


def _clean_text(v):
    if v is None:
        return None
    if isinstance(v, float):
        if math.isnan(v):
            return None
        if v.is_integer():
            return str(int(v))
        return str(v)
    if isinstance(v, int):
        return str(v)
    if isinstance(v, Decimal):
        return str(v)
    s = str(v).strip()
    return s or None


def _read_rows(content: bytes, filename: str):
    name = (filename or "").lower()
    if name.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="ignore")
        reader = csv.reader(io.StringIO(text))
        rows = [list(r) for r in reader]
    elif name.endswith(".xlsx") or name.endswith(".xls"):
        wb = load_workbook(BytesIO(content), data_only=True)
        ws = wb.active
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        raise ValueError("仅支持 .xlsx / .xls / .csv 格式的文件")

    if not rows:
        raise ValueError("文件内容为空，请使用模板填写后再上传")

    header = [_clean_text(h) for h in rows[0]]
    data_rows = rows[1:]
    return header, data_rows


def import_scores_from_file(content: bytes, filename: str, db: Session) -> dict:
    header, data_rows = _read_rows(content, filename)

    header_set = set(h for h in header if h)
    missing_cols = [col for col in IMPORT_HEADERS if col not in header_set]
    if missing_cols:
        raise ValueError(f"模板表头缺少必填列：{'、'.join(missing_cols)}，请下载标准模板填写")

    col_index = {IMPORT_COLUMN_MAP[h]: idx for idx, h in enumerate(header) if h in IMPORT_COLUMN_MAP}

    valid_infos: List[Addscore] = []
    failures: list[dict] = []
    effective_total = 0
    seen_pairs: set[tuple[str, int]] = set()

    for offset, row in enumerate(data_rows):
        excel_row_no = offset + 2

        if row is None or all(_clean_text(c) is None for c in row):
            continue
        effective_total += 1

        def cell(field):
            idx = col_index.get(field)
            if idx is None or idx >= len(row):
                return None
            return _clean_text(row[idx])

        student_no = cell("student_no")
        try:
            exam_order_raw = cell("exam_order")
            if exam_order_raw is None:
                exam_order = None
            else:
                try:
                    exam_order = int(float(exam_order_raw))
                except (ValueError, TypeError):
                    raise ValueError(f"考试次序 必须是数字（当前填的是「{exam_order_raw}」）")

            score_raw = cell("score")
            if score_raw is None:
                score = None
            else:
                try:
                    score = Decimal(str(score_raw))
                except (InvalidOperation, ValueError, TypeError):
                    raise ValueError(f"成绩 必须是 0~100 的数字（当前填的是「{score_raw}」）")

            payload = {
                "student_no": student_no,
                "exam_order": exam_order,
                "score": score,
            }
            info = Addscore(**payload)

            pair = (info.student_no, info.exam_order)
            if pair in seen_pairs:
                raise ValueError(f"同一文件中学号 {info.student_no} 的第{info.exam_order} 次成绩重复")
            seen_pairs.add(pair)

            student = get_by_student_no(info.student_no, db)
            if not student:
                raise ValueError(f"学生 {info.student_no} 的信息不在学生表中")

            if score_dao.is_exist(db, info.student_no, info.exam_order):
                raise ValueError(f"学生 {info.student_no} 第{info.exam_order} 次成绩已存在")

            valid_infos.append(info)
        except ValidationError as ve:
            reasons = []
            for err in ve.errors():
                loc = err.get("loc") or []
                field = loc[0] if loc else ""
                reasons.append(f"{FIELD_CN.get(field, field)}：{err.get('msg')}")
            failures.append({"row": excel_row_no, "student_no": student_no, "reason": "；".join(reasons)})
        except (ValueError, TypeError) as e:
            failures.append({"row": excel_row_no, "student_no": student_no, "reason": str(e)})

    if valid_infos:
        batch_data = [
            Score(student_no=info.student_no, exam_order=info.exam_order, score=info.score)
            for info in valid_infos
        ]
        score_dao.batch_add_scores(db, batch_data)

    result = {
        "total": effective_total,
        "success_count": len(valid_infos),
        "fail_count": len(failures),
        "failures": failures,
    }
    logger.info("成绩导入完成：共 %s 行，成功 %s 条，失败 %s 条",
                effective_total, len(valid_infos), len(failures))
    return result


def build_import_template() -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "成绩导入模板"

    ws.append(IMPORT_HEADERS)
    ws.append(["S2025001", 1, 90])
    ws.append(["S2025002", 1, 75])

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2563EB")
    center = Alignment(horizontal="center", vertical="center")
    widths = {1: 18, 2: 14, 3: 12}

    for col_idx, _ in enumerate(IMPORT_HEADERS, start=1):
        c = ws.cell(row=1, column=col_idx)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        ws.column_dimensions[c.column_letter].width = widths.get(col_idx, 16)

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def update_score(data: Updatescore, db: Session):
    """更新学生成绩（含业务校验）"""
    is_exist = score_dao.is_exist(db, data.student_no, data.exam_order)
    if not is_exist:
        logger.warning("修改成绩校验失败：学生 %s 第%s次成绩不存在",
                       data.student_no, data.exam_order)
        raise ValueError("要修改的学生信息不存在")
    return score_dao.update_score(db, data.student_no, data.exam_order, data.score)


def soft_delete_score(student_no: str, exam_order: int, db: Session):
    """逻辑删除学生成绩（含业务校验）"""
    is_exist = score_dao.is_exist(db, student_no, exam_order)
    if not is_exist:
        logger.warning("删除成绩校验失败：学生 %s 第%s次成绩不存在", student_no, exam_order)
        raise ValueError("要删除的学生信息不存在")
    return score_dao.is_delete(db, student_no, exam_order)


def query_scores(
    db: Session,
    page: int = 1,
    page_size: int = 10,
    student_no: str = None,
    exam_order: int = None,
    class_id: int = None,
    min_score: Decimal = None,
    max_score: Decimal = None,
    sort_no: str = None,
    sort_score: str = None,
):
    """查询学生成绩（含业务校验）"""
    if min_score is not None and max_score is not None:
        if min_score > max_score:
            logger.warning("查询成绩校验失败：最小范围 %s 大于最大范围 %s", min_score, max_score)
            raise ValueError("请求异常,最小范围不能大于最大范围")

    result = score_dao.query_score(
        db, page, page_size,
        student_no, exam_order, class_id,
        min_score, max_score,
        sort_no, sort_score
    )

    if not result:
        logger.warning("查询成绩：无符合条件的数据，student_no=%s, exam_order=%s, class_id=%s",
                       student_no, exam_order, class_id)
        raise ValueError("要查询的学生信息不存在")

    total = score_dao.count_score(
        db, student_no, exam_order, class_id, min_score, max_score
    )

    return result, page, page_size, total
