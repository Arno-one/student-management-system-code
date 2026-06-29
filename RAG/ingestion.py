"""
文档解析、切片、Embedding、入库的完整离线管线。

支持两类知识库：
1. novels: 四大名著 txt + QA + 通用 Markdown
2. production: PDF / DOCX / Markdown / QA
"""
from __future__ import annotations

import json
import re
import time
from pathlib import Path

from langchain_text_splitters import RecursiveCharacterTextSplitter

from RAG.clients import embed_texts, get_milvus_client
from RAG.config import config
from RAG.schemas import get_or_create_collection, load_collection
from util.log import get_logger

logger = get_logger(__name__)

_CHAPTER_RE = re.compile(r"^第[一二三四五六七八九十百千零\d]+回")
_MD_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+)$", re.MULTILINE)
_CN_SECTION_RE = re.compile(r"^[一二三四五六七八九十]+、.+$")
_NUM_SECTION_RE = re.compile(r"^\d+(?:\.\d+)*[\s、].+$")


def _find_novel_files(docs_dir: Path) -> list[Path]:
    """扫描四大名著 txt 文件。"""
    files = sorted(docs_dir.glob("*.txt"))
    if not files:
        raise FileNotFoundError(f"未在 {docs_dir} 找到任何 txt 文件")
    logger.info("发现 %d 个名著 txt 文件: %s", len(files), [file.name for file in files])
    return files


def _find_qa_files(docs_dir: Path) -> list[Path]:
    """扫描 Q&A 文件，命名规则沿用 *Q&A*.md。"""
    files = sorted(docs_dir.rglob("*Q&A*.md"))
    logger.info("发现 %d 个 QA 文件: %s", len(files), [str(file.relative_to(docs_dir)) for file in files])
    return files


def _find_generic_files(docs_dir: Path) -> list[Path]:
    """按当前知识库配置扫描除 QA 外的普通文档。"""
    qa_names = {file.resolve() for file in _find_qa_files(docs_dir)}
    files = []
    for ext in config.current_kb.supported_extensions:
        for file in sorted(docs_dir.rglob(f"*{ext}")):
            if file.resolve() in qa_names:
                continue
            files.append(file)

    unique_files = sorted({file.resolve(): file for file in files}.values(), key=lambda item: str(item))
    if unique_files:
        logger.info(
            "发现 %d 个通用文档: %s",
            len(unique_files),
            [str(file.relative_to(docs_dir)) for file in unique_files],
        )
    return unique_files


def _parse_qa_file(filepath: Path) -> list[dict]:
    """解析现有约定的 QA markdown 文件：标题 + JSON 数组。"""
    text = filepath.read_text(encoding="utf-8")
    text = re.sub(r"^# .*\r?\n", "", text)
    text = text.replace("\\[", "[").replace("\\]", "]")
    text = text.replace("\\{", "{").replace("\\}", "}")
    text = text.replace("\\&", "&")
    text = re.sub(r'\\(?!["\\/bfnrtu])', "", text)
    text = text.strip()

    last_bracket = text.rfind("]")
    if last_bracket >= 0:
        text = text[: last_bracket + 1]

    pairs = json.loads(text)
    valid = [item for item in pairs if all(item.get(key) for key in ("question", "answer", "reason"))]
    if len(valid) != len(pairs):
        logger.warning("%s: 仅 %d/%d 条 QA 满足 question/answer/reason 要求", filepath.name, len(valid), len(pairs))
    return valid


def _parse_chapters(text: str, max_chapters: int | None) -> list[tuple[str, int]]:
    """按四大名著“第 X 回”格式切分章节。"""
    lines = text.splitlines()
    chapters: list[tuple[str, int]] = []
    current_lines: list[str] = []
    current_chapter_num = 0
    in_preamble = True

    for raw_line in lines:
        line = raw_line.strip()
        if _CHAPTER_RE.match(line):
            if current_lines and current_chapter_num > 0:
                chapters.append(("\n".join(current_lines), current_chapter_num))
            current_chapter_num += 1
            if max_chapters and current_chapter_num > max_chapters:
                break
            current_lines = [line]
            in_preamble = False
            continue

        if not in_preamble:
            current_lines.append(line)

    if current_lines and current_chapter_num > 0 and (not max_chapters or current_chapter_num <= max_chapters):
        chapters.append(("\n".join(current_lines), current_chapter_num))
    return chapters


def _clean_text(text: str) -> str:
    """基础清洗：去首尾空白、压缩多余空行与空格。"""
    cleaned = text.strip()
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    cleaned = re.sub(r" {3,}", "  ", cleaned)
    return cleaned


def _split_chunks(text: str) -> list[str]:
    """按统一配置做字符级切片。"""
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=config.chunk_size,
        chunk_overlap=config.chunk_overlap,
        separators=["\n\n", "\n", "。", "；", "，", " ", ""],
    )
    return splitter.split_text(text)


def _looks_like_heading(line: str) -> bool:
    """判断一行文本是否像结构化标题。"""
    stripped = line.strip()
    if not stripped:
        return False
    if len(stripped) > 50:
        return False
    return bool(_CN_SECTION_RE.match(stripped) or _NUM_SECTION_RE.match(stripped))


def _parse_markdown_sections(text: str) -> list[tuple[str, str]]:
    """按 Markdown 标题切分。没有标题时回退全文。"""
    matches = list(_MD_HEADING_RE.finditer(text))
    if not matches:
        return [("全文", text.strip())] if text.strip() else []

    sections = []
    preamble = text[: matches[0].start()].strip()
    if preamble:
        sections.append(("前言", preamble))

    for index, match in enumerate(matches):
        title = match.group(2).strip()
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        body = text[start:end].strip()
        if body:
            sections.append((title, body))
    return sections


def _parse_plain_sections(text: str) -> list[tuple[str, str]]:
    """按中文/数字标题行切分普通文本，切不出来就整篇回退。"""
    lines = [line.rstrip() for line in text.splitlines()]
    sections: list[tuple[str, str]] = []
    current_title = "全文"
    current_lines: list[str] = []
    saw_heading = False

    for line in lines:
        stripped = line.strip()
        if _looks_like_heading(stripped):
            saw_heading = True
            if current_lines:
                body = _clean_text("\n".join(current_lines))
                if body:
                    sections.append((current_title, body))
            current_title = stripped
            current_lines = []
            continue

        current_lines.append(line)

    if current_lines:
        body = _clean_text("\n".join(current_lines))
        if body:
            sections.append((current_title, body))

    if saw_heading and sections:
        return sections

    whole = _clean_text(text)
    return [("全文", whole)] if whole else []


def _extract_pdf_sections(filepath: Path) -> list[tuple[str, str]]:
    """提取 PDF 文本，并优先按章节标题切分。"""
    from pypdf import PdfReader

    reader = PdfReader(str(filepath))
    pages = []
    for page_index, page in enumerate(reader.pages, start=1):
        page_text = (page.extract_text() or "").strip()
        if not page_text:
            continue
        pages.append(f"第{page_index}页\n{page_text}")

    raw_text = "\n\n".join(pages)
    sections = _parse_plain_sections(raw_text)
    return sections


def _extract_docx_sections(filepath: Path) -> list[tuple[str, str]]:
    """提取 DOCX 文本，并优先按标题样式切分。"""
    from docx import Document

    document = Document(str(filepath))
    sections: list[tuple[str, str]] = []
    current_title = "全文"
    current_lines: list[str] = []
    saw_heading = False

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        style_name = getattr(paragraph.style, "name", "") or ""
        is_heading = (
            "Heading" in style_name
            or "标题" in style_name
            or _looks_like_heading(text)
        )

        if is_heading:
            saw_heading = True
            if current_lines:
                body = _clean_text("\n".join(current_lines))
                if body:
                    sections.append((current_title, body))
            current_title = text
            current_lines = []
            continue

        current_lines.append(text)

    if current_lines:
        body = _clean_text("\n".join(current_lines))
        if body:
            sections.append((current_title, body))

    if saw_heading and sections:
        return sections

    whole_text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
    return _parse_plain_sections(whole_text)


def _extract_markdown_sections(filepath: Path) -> list[tuple[str, str]]:
    """提取 Markdown 文本，并按标题切片。"""
    text = filepath.read_text(encoding="utf-8")
    text = re.sub(r"^# .*\r?\n", "", text)
    return _parse_markdown_sections(text)


def _extract_generic_text_sections(filepath: Path) -> list[tuple[str, str]]:
    """兜底文本解析。"""
    text = filepath.read_text(encoding="utf-8")
    return _parse_plain_sections(text)


def _extract_document_sections(filepath: Path) -> tuple[str, list[tuple[str, str]]]:
    """按文件类型路由到对应的解析器。"""
    suffix = filepath.suffix.lower()
    if suffix == ".pdf":
        return "pdf", _extract_pdf_sections(filepath)
    if suffix == ".docx":
        return "docx", _extract_docx_sections(filepath)
    if suffix == ".md":
        return "md", _extract_markdown_sections(filepath)
    if suffix == ".txt":
        return "txt", _extract_generic_text_sections(filepath)
    raise ValueError(f"暂不支持的文档类型: {filepath}")


def _build_chunks_from_sections(
    filepath: Path,
    docs_dir: Path,
    source_type: str,
    sections: list[tuple[str, str]],
) -> tuple[list[dict], dict]:
    """把解析出的章节内容统一转成待入库 chunk。"""
    rel_path = str(filepath.relative_to(docs_dir))
    chunks: list[dict] = []
    total_sections = 0

    for section_index, (section_title, section_text) in enumerate(sections):
        cleaned = _clean_text(section_text)
        if not cleaned:
            continue

        total_sections += 1
        split_result = _split_chunks(cleaned)
        total = len(split_result)
        section_doc_id = f"{filepath.stem}_sec_{section_index:03d}"
        for index, chunk_text in enumerate(split_result):
            enriched_text = f"【{section_title}】{chunk_text}" if section_title else chunk_text
            chunks.append(
                {
                    "doc_id": section_doc_id,
                    "text": enriched_text,
                    "question": "",
                    "answer": "",
                    "reason": "",
                    "file_name": rel_path,
                    "chunk_index": index,
                    "total_chunks": total,
                    "source_type": source_type,
                    "created_at": int(time.time()),
                }
            )

    stats = {
        "file": rel_path,
        "source_type": source_type,
        "sections": total_sections,
        "chunks": len(chunks),
    }
    return chunks, stats


def ingest_documents(drop_existing: bool = False) -> dict:
    """
    完整入库管线：读取文档 -> 清洗 -> 切片 -> Embedding -> 写入 Milvus。
    """
    t_start = time.time()
    docs_dir = config.docs_dir
    collection = get_or_create_collection(drop_existing=drop_existing)

    all_chunks: list[dict] = []
    novel_stats = []
    qa_stats = []
    doc_stats = []

    if config.active_kb_id == "novels":
        max_chapters = config.max_chapters
        for filepath in _find_novel_files(docs_dir):
            file_t0 = time.time()
            raw_text = filepath.read_text(encoding="utf-8")
            chapters = _parse_chapters(raw_text, max_chapters)
            if not chapters:
                logger.warning("%s: 未解析到任何章回，跳过", filepath.name)
                continue

            chunk_count = 0
            for chapter_text, chapter_num in chapters:
                split_result = _split_chunks(_clean_text(chapter_text))
                total = len(split_result)
                for index, chunk_text in enumerate(split_result):
                    all_chunks.append(
                        {
                            "doc_id": f"{filepath.stem}_{chapter_num:03d}",
                            "text": chunk_text,
                            "question": "",
                            "answer": "",
                            "reason": "",
                            "file_name": filepath.name,
                            "chunk_index": index,
                            "total_chunks": total,
                            "source_type": "txt",
                            "created_at": int(time.time()),
                        }
                    )
                    chunk_count += 1

            logger.info("%s: %d 回 -> %d chunks (%.1fs)", filepath.name, len(chapters), chunk_count, time.time() - file_t0)
            novel_stats.append({"file": filepath.name, "chapters": len(chapters), "chunks": chunk_count})

    for filepath in _find_qa_files(docs_dir):
        file_t0 = time.time()
        pairs = _parse_qa_file(filepath)
        if not pairs:
            logger.warning("%s: 无有效 QA 对，跳过", filepath.name)
            continue

        total = len(pairs)
        for index, pair in enumerate(pairs):
            search_text = f"{pair['question']} {pair['answer']} {pair['reason']}"
            all_chunks.append(
                {
                    "doc_id": f"{filepath.stem}_qa_{index:03d}",
                    "text": search_text,
                    "question": pair["question"],
                    "answer": pair["answer"],
                    "reason": pair["reason"],
                    "file_name": str(filepath.relative_to(docs_dir)),
                    "chunk_index": index,
                    "total_chunks": total,
                    "source_type": "qa",
                    "created_at": int(time.time()),
                }
            )

        logger.info("%s: %d 对 QA -> %d 条 (%.1fs)", filepath.name, total, total, time.time() - file_t0)
        qa_stats.append({"file": str(filepath.relative_to(docs_dir)), "qa_pairs": total})

    for filepath in _find_generic_files(docs_dir):
        file_t0 = time.time()
        source_type, sections = _extract_document_sections(filepath)
        if not sections:
            logger.warning("%s: 未提取到有效内容，跳过", filepath.name)
            continue

        chunks, stats = _build_chunks_from_sections(filepath, docs_dir, source_type, sections)
        if not chunks:
            logger.warning("%s: 未生成任何 chunk，跳过", filepath.name)
            continue

        all_chunks.extend(chunks)
        logger.info(
            "%s: %d 节 -> %d chunks (%.1fs)",
            stats["file"],
            stats["sections"],
            stats["chunks"],
            time.time() - file_t0,
        )
        doc_stats.append(stats)

    if not all_chunks:
        logger.warning("[%s] 没有生成任何 chunk，入库终止", config.active_kb_id)
        return {
            "total_chunks": 0,
            "files": novel_stats,
            "qa_files": qa_stats,
            "doc_files": doc_stats,
            "elapsed_seconds": round(time.time() - t_start, 2),
        }

    source_counts: dict[str, int] = {}
    for chunk in all_chunks:
        source_type = chunk["source_type"]
        source_counts[source_type] = source_counts.get(source_type, 0) + 1

    logger.info("[%s] 共生成 %d 个条目，开始批量 Embedding", config.active_kb_id, len(all_chunks))
    vectors = embed_texts([chunk["text"] for chunk in all_chunks])
    if len(vectors) != len(all_chunks):
        raise RuntimeError(f"Embedding 返回数量 ({len(vectors)}) 与 chunk 数量 ({len(all_chunks)}) 不一致")

    insert_data = []
    for chunk, vector in zip(all_chunks, vectors):
        if len(vector) != config.embedding_dimension:
            raise RuntimeError(f"向量维度 {len(vector)} 与配置 {config.embedding_dimension} 不一致")
        insert_data.append(
            {
                "doc_id": chunk["doc_id"],
                "text": chunk["text"],
                "question": chunk["question"],
                "answer": chunk["answer"],
                "reason": chunk["reason"],
                "file_name": chunk["file_name"],
                "chunk_index": chunk["chunk_index"],
                "total_chunks": chunk["total_chunks"],
                "source_type": chunk["source_type"],
                "created_at": chunk["created_at"],
                "dense_vector": vector,
            }
        )

    client = get_milvus_client()
    batch_size = 500
    for start in range(0, len(insert_data), batch_size):
        batch = insert_data[start : start + batch_size]
        client.insert(collection_name=config.milvus_collection_name, data=batch)
        logger.debug(
            "[%s] 插入 batch %d/%d",
            config.active_kb_id,
            start // batch_size + 1,
            (len(insert_data) + batch_size - 1) // batch_size,
        )

    client.flush(config.milvus_collection_name)
    load_collection(collection)

    total_elapsed = round(time.time() - t_start, 2)
    logger.info(
        "[%s] 入库完成: total=%d, source_counts=%s, %.1fs",
        config.active_kb_id,
        len(insert_data),
        source_counts,
        total_elapsed,
    )
    return {
        "total_chunks": len(insert_data),
        "files": novel_stats,
        "qa_files": qa_stats,
        "doc_files": doc_stats,
        "source_counts": source_counts,
        "elapsed_seconds": total_elapsed,
        "collection_name": config.milvus_collection_name,
        "kb_id": config.active_kb_id,
    }
